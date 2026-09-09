from __future__ import annotations

from pathlib import Path
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path("/Users/vincent/GeekVincent/Codex/CreatorOS")
OUT = ROOT / "outputs" / "creatoros-ai-pet-plus-business-plan.docx"
HERO = ROOT / "assets" / "pitch" / "ai-pet-studio-hero.png"
MATRIX = ROOT / "assets" / "pitch" / "ai-pet-ip-matrix.png"

INK = RGBColor(20, 28, 33)
BLUE = RGBColor(46, 116, 181)
DARK_BLUE = RGBColor(31, 77, 120)
MUTED = RGBColor(91, 103, 112)
ORANGE = RGBColor(242, 106, 46)
TABLE_FILL = "F4F6F9"
LIGHT_FILL = "F8F6F2"
WARN_FILL = "FFF1E8"
BORDER = "D9DEE6"
FONT_LATIN = "Calibri"
FONT_CJK = "Microsoft YaHei"


def set_run_font(run, size=None, color=None, bold=None, italic=None):
    run.font.name = FONT_LATIN
    run._element.rPr.rFonts.set(qn("w:ascii"), FONT_LATIN)
    run._element.rPr.rFonts.set(qn("w:hAnsi"), FONT_LATIN)
    run._element.rPr.rFonts.set(qn("w:eastAsia"), FONT_CJK)
    if size is not None:
        run.font.size = Pt(size)
    if color is not None:
        run.font.color.rgb = color
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def set_paragraph_spacing(paragraph, before=0, after=8, line=1.333):
    paragraph.paragraph_format.space_before = Pt(before)
    paragraph.paragraph_format.space_after = Pt(after)
    paragraph.paragraph_format.line_spacing = line


def shade_cell(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, v in {"top": top, "start": start, "bottom": bottom, "end": end}.items():
        node = tc_mar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")


def set_cell_text(cell, text, bold=False, size=10.5, color=INK, align=None):
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    cell.text = ""
    p = cell.paragraphs[0]
    if align is not None:
        p.alignment = align
    set_paragraph_spacing(p, after=0, line=1.15)
    run = p.add_run(text)
    set_run_font(run, size=size, color=color, bold=bold)
    set_cell_margins(cell)


def set_table_borders(table, color=BORDER):
    tbl = table._tbl
    tbl_pr = tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = f"w:{edge}"
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), "4")
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), color)


def set_table_width(table, widths_inches):
    table.autofit = False
    for row in table.rows:
        for idx, width in enumerate(widths_inches):
            row.cells[idx].width = Inches(width)


def add_para(doc, text="", style=None, size=11, color=INK, bold=False, italic=False, align=None, after=8, before=0):
    p = doc.add_paragraph(style=style)
    if align is not None:
        p.alignment = align
    set_paragraph_spacing(p, before=before, after=after)
    if text:
        run = p.add_run(text)
        set_run_font(run, size=size, color=color, bold=bold, italic=italic)
    return p


def add_heading(doc, text, level=1):
    style = f"Heading {level}"
    p = doc.add_paragraph(style=style)
    run = p.add_run(text)
    if level == 1:
        set_run_font(run, size=16, color=BLUE, bold=True)
        set_paragraph_spacing(p, before=18, after=10)
    elif level == 2:
        set_run_font(run, size=13, color=BLUE, bold=True)
        set_paragraph_spacing(p, before=12, after=6)
    else:
        set_run_font(run, size=12, color=DARK_BLUE, bold=True)
        set_paragraph_spacing(p, before=8, after=4)
    return p


def add_bullet(doc, text):
    p = doc.add_paragraph(style="List Bullet")
    set_paragraph_spacing(p, after=4, line=1.208)
    run = p.add_run(text)
    set_run_font(run, size=11, color=INK)
    return p


def add_number(doc, text):
    p = doc.add_paragraph(style="List Number")
    set_paragraph_spacing(p, after=4, line=1.208)
    run = p.add_run(text)
    set_run_font(run, size=11, color=INK)
    return p


def add_callout(doc, title, body, fill=LIGHT_FILL):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    table.columns[0].width = Inches(6.35)
    set_table_borders(table, color="E4E7EC")
    cell = table.cell(0, 0)
    shade_cell(cell, fill)
    set_cell_margins(cell, top=120, bottom=120, start=160, end=160)
    cell.text = ""
    p1 = cell.paragraphs[0]
    set_paragraph_spacing(p1, after=3, line=1.15)
    r1 = p1.add_run(title)
    set_run_font(r1, size=11.5, color=ORANGE, bold=True)
    p2 = cell.add_paragraph()
    set_paragraph_spacing(p2, after=0, line=1.25)
    r2 = p2.add_run(body)
    set_run_font(r2, size=10.5, color=INK)
    add_para(doc, "", after=4)


def add_table(doc, headers, rows, widths):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    set_table_borders(table)
    set_table_width(table, widths)
    for idx, header in enumerate(headers):
        cell = table.rows[0].cells[idx]
        shade_cell(cell, TABLE_FILL)
        set_cell_text(cell, header, bold=True, size=10.5, color=INK)
    for row_values in rows:
        row = table.add_row()
        for idx, value in enumerate(row_values):
            cell = row.cells[idx]
            set_cell_text(cell, value, size=10.2, color=INK)
    set_table_width(table, widths)
    add_para(doc, "", after=4)
    return table


def configure_document(doc):
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = FONT_LATIN
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), FONT_CJK)
    normal.font.size = Pt(11)
    normal.font.color.rgb = INK
    normal.paragraph_format.space_after = Pt(8)
    normal.paragraph_format.line_spacing = 1.333

    for style_name in ["Heading 1", "Heading 2", "Heading 3"]:
        style = styles[style_name]
        style.font.name = FONT_LATIN
        style._element.rPr.rFonts.set(qn("w:eastAsia"), FONT_CJK)
        style.font.bold = True

    header = section.header.paragraphs[0]
    header.text = ""
    header.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run = header.add_run("CreatorOS AI 萌宠+ 内容工厂 | Business Plan")
    set_run_font(run, size=9, color=MUTED)

    footer = section.footer.paragraphs[0]
    footer.text = ""
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = footer.add_run("Confidential draft for discussion | 数据与融资假设需在正式路演前复核")
    set_run_font(run, size=8.5, color=MUTED)


def build_doc():
    doc = Document()
    configure_document(doc)

    # Cover
    add_para(doc, "CreatorOS", size=12, color=MUTED, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, after=8)
    add_para(doc, "AI 萌宠+ 内容工厂", size=26, color=INK, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, after=4)
    add_para(
        doc,
        "业务计划书",
        size=16,
        color=ORANGE,
        bold=True,
        align=WD_ALIGN_PARAGRAPH.CENTER,
        after=10,
    )
    add_para(
        doc,
        "用 AI Agent 批量孵化萌宠垂类 IP，先以内容矩阵验证流量，再沉淀为可规模化的创作者增长系统。",
        size=12,
        color=MUTED,
        align=WD_ALIGN_PARAGRAPH.CENTER,
        after=10,
    )
    meta = [
        ("版本", "V0.2 业务计划书草案"),
        ("日期", "2026 年 7 月 1 日"),
        ("融资阶段", "建议天使轮 / Angel Round"),
        ("建议融资额", "人民币 500 万，可按路演对象调整"),
    ]
    table = doc.add_table(rows=len(meta), cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(table, color="E4E7EC")
    set_table_width(table, [1.55, 4.35])
    for row, (label, value) in zip(table.rows, meta):
        shade_cell(row.cells[0], TABLE_FILL)
        set_cell_text(row.cells[0], label, bold=True, size=10.2)
        set_cell_text(row.cells[1], value, size=10.2)
    add_para(doc, "", after=4)
    if HERO.exists():
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run()
        r.add_picture(str(HERO), width=Inches(5.6))
    add_para(
        doc,
        "说明：本计划书根据 CreatorOS AI 萌宠+ 内容工厂融资路演 PPT、CreatorOS PRD 和当前 MVP 功能整理；市场数据为公开资料与路演稿引用口径，正式融资前需再次核验。",
        size=9.5,
        color=MUTED,
        italic=True,
        align=WD_ALIGN_PARAGRAPH.CENTER,
        after=0,
    )
    doc.add_page_break()

    add_heading(doc, "一、执行摘要", 1)
    add_para(
        doc,
        "CreatorOS AI 萌宠+ 内容工厂是一套“内容矩阵 + AI Agent + 创作者增长系统”的双引擎业务。短期以萌宠+脱口秀、萌宠+教育、萌宠+职场、萌宠+疗愈等账号矩阵验证流量和商业化；中长期把经过真实账号验证的选题、爆款逆向、角色设定、脚本生成、发布复盘能力产品化，形成面向创作者、MCN、品牌内容团队的 SaaS / Agent 平台。",
    )
    add_callout(
        doc,
        "核心判断",
        "我们不是再做一个单点 AI 视频工具，而是做“可复制账号增长方法 + AI 角色内容生产 + 数据反馈”的闭环。内容账号既是获客渠道，也是产品训练场和案例资产。",
        fill=WARN_FILL,
    )
    add_table(
        doc,
        ["维度", "计划书口径"],
        [
            ("项目定位", "面向自媒体创作者与内容团队的 AI 创作者增长系统，并自孵化 AI 萌宠+ 内容 IP。"),
            ("首个内容切口", "萌宠+，因其具备强情绪价值、跨人群传播、角色可持续和商业化延展空间。"),
            ("核心产品", "趋势雷达、爆款逆向器、角色圣经、内容 Agent、审批发布、复盘回流。"),
            ("12 个月目标", "10 个种子账号、1000 条内容测试、SaaS MVP 收费试点。"),
            ("融资建议", "天使轮人民币 500 万，用于产品研发、账号矩阵运营、团队建设、数据/版权合规和商务拓展。"),
        ],
        [1.35, 5.0],
    )

    add_heading(doc, "二、市场机会", 1)
    add_para(
        doc,
        "本项目的窗口来自三类趋势交汇：第一，宠物与萌宠内容具有天然情绪价值，能承接陪伴、解压、亲子、教育、职场吐槽等多场景内容需求；第二，创作者经济仍在扩张，但普通创作者面临起号难、制作慢、复盘弱的问题；第三，生成式 AI 正在把脚本、图像、视频和多模态内容生产成本持续降低，使“账号矩阵化”和“角色资产化”成为可操作路径。",
    )
    add_table(
        doc,
        ["趋势", "对本项目的意义", "计划书引用口径"],
        [
            ("宠物情绪消费", "萌宠内容可承载陪伴、治愈、吐槽、教育等多个内容场景，适合建立长期 IP。", "中国宠物消费支出规模约 ¥300B，且仍在增长。"),
            ("创作者经济", "创作者需要更稳定的选题、定位、内容生产和复盘工具，B 端产品化空间存在。", "Goldman Sachs 预测创作者经济 2027 年接近 $480B。"),
            ("AI 内容生产", "AI 可降低脚本、分镜、封面、视频生成和平台适配成本，让小团队具备矩阵化能力。", "Adobe 调研口径：全球创作者中已有较高比例使用生成式 AI。"),
            ("AI 陪伴/玩具", "萌宠角色可以从内容 IP 延展到互动陪伴、玩具、周边和品牌合作。", "公开报道提到中国 AI 玩具/陪伴宠物市场 2030 年预测约 $1.4B+。"),
        ],
        [1.25, 3.0, 2.1],
    )
    add_para(
        doc,
        "以上数字用于说明赛道方向和投资逻辑，不构成财务预测。正式融资版本建议补充来自权威行业报告、平台数据和自有账号测试数据的可核验证据。",
        size=10,
        color=MUTED,
        italic=True,
    )

    add_heading(doc, "三、用户痛点与机会缺口", 1)
    add_para(
        doc,
        "目标用户包括短期起号创作者、长期经营创作者、小型 MCN 和品牌内容团队。当前内容生产链路中，创作者并不缺少单个生成工具，而是缺少一套从“发现机会”到“生成内容”再到“发布复盘”的系统化方法。",
    )
    add_bullet(doc, "起号靠感觉：创作者不知道该追哪类热点，也难判断爆款结构是否适合自己的账号定位。")
    add_bullet(doc, "制作链条长：角色设定、脚本、分镜、图文、视频、多平台文案都需要反复手工处理。")
    add_bullet(doc, "内容同质化：只模仿表层钩子，无法形成稳定人设、栏目资产和可持续复利。")
    add_bullet(doc, "反馈断裂：发布、数据回流、复盘和下一轮选题之间缺少自动闭环。")
    add_callout(
        doc,
        "要解决的核心问题",
        "不是“生成一条内容”，而是让一批账号持续找到可复制的增长动作，并把有效动作沉淀成产品能力。",
    )

    add_heading(doc, "四、解决方案：CreatorOS 是萌宠 IP 工厂的操作系统", 1)
    add_para(
        doc,
        "CreatorOS 将内容生产拆成六个可产品化环节：趋势发现、爆款逆向、角色设定、内容生成、审批发布、复盘回流。内容团队可以先在自有 AI 萌宠账号矩阵中使用这套系统，再把有效流程外部化为 SaaS / Agent 产品。",
    )
    add_table(
        doc,
        ["模块", "功能说明", "对业务的价值"],
        [
            ("趋势雷达", "监控平台热度、增速、内容形式和可复制度。", "帮助短期起号创作者快速找到可复制信号。"),
            ("爆款逆向器", "根据热点视频链接、标题、字幕/观察笔记，输出结构拆解与 prompt 包。", "把爆款从“灵感”变成可复用生产配方，同时保留改写边界。"),
            ("角色圣经", "维护每个萌宠 IP 的人设、口癖、价值观、禁区和栏目方向。", "保证账号矩阵在批量生产时保持人格一致。"),
            ("内容 Agent", "生成标题、口播稿、图文页、视频分镜、标签和 CTA。", "降低内容生产成本，提高多平台适配效率。"),
            ("审批发布", "人工确认后进入发布队列，未来接入平台授权发布。", "保留内容安全和合规审核。"),
            ("复盘模型", "根据表现数据调整下一轮选题、脚本和发布时间。", "让账号运营经验形成可学习的数据资产。"),
        ],
        [1.25, 2.7, 2.4],
    )

    add_heading(doc, "五、内容业务：AI 萌宠+ 账号矩阵", 1)
    add_para(
        doc,
        "“萌宠+”不是单个账号，而是一组垂类 IP 组合。它用同一套底层生产系统，孵化不同内容场景和商业化路径。每个账号都需要固定角色、固定栏目、固定转化路径，并在复盘中持续沉淀角色资产和爆款结构。",
    )
    if MATRIX.exists():
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.add_run().add_picture(str(MATRIX), width=Inches(6.3))
        add_para(doc, "图：AI 萌宠+ IP 矩阵示意图", size=9.5, color=MUTED, italic=True, align=WD_ALIGN_PARAGRAPH.CENTER)
    add_table(
        doc,
        ["账号方向", "目标人群", "内容钩子", "商业化路径"],
        [
            ("萌宠+脱口秀", "泛娱乐用户", "宠物替人类吐槽热点，形成高频短视频栏目。", "品牌植入、IP 授权、直播互动。"),
            ("萌宠+教育", "亲子与学生", "用萌宠角色讲英语、科学、常识和学习方法。", "课程联名、教育广告、内容授权。"),
            ("萌宠+职场", "打工人", "宠物演绎办公室荒诞和职场沟通痛点。", "招聘平台、办公 SaaS、品牌合作。"),
            ("萌宠+疗愈", "年轻女性与高压人群", "睡前陪伴、冥想、心理陪伴、轻 ASMR。", "音频会员、周边、电商转化。"),
            ("萌宠+财经小白", "理财入门用户", "用宠物故事讲基础财商概念。", "金融合规广告、内容课程。"),
            ("萌宠+生活方式", "养宠家庭", "剧情种草、宠物用品测评、家居旅行场景。", "电商佣金、自营商品、品牌共创。"),
        ],
        [1.25, 1.3, 2.25, 1.55],
    )

    add_heading(doc, "六、增长飞轮", 1)
    add_para(
        doc,
        "项目的关键不是单次内容生产效率，而是“内容表现越多，Agent 越懂赛道；Agent 越准，账号孵化越快”。因此增长飞轮从热点信号开始，经角色适配、批量生产、多平台发布、数据复盘，最后反哺产品强化。",
    )
    for step in [
        "热点信号：从平台内容数据中找到正在起量的结构。",
        "角色适配：判断信号是否适合某个萌宠 IP 的人设和栏目定位。",
        "批量生产：Agent 生成多版本脚本、封面、分镜和发布文案。",
        "多平台发布：人工确认后分发到目标平台。",
        "数据复盘：沉淀有效钩子、角色反应、转化路径和评论问题。",
        "产品强化：把高胜率流程固化为 CreatorOS 的标准能力。",
    ]:
        add_number(doc, step)

    add_heading(doc, "七、产品化路径与技术路线", 1)
    add_para(
        doc,
        "当前 MVP 已具备本地工作台能力：趋势雷达、爆款逆向器、定位策略、内容 Agent、审批发布队列和本地 JSON 持久化。下一阶段需要把模拟数据和规则生成升级为真实数据接入、模型调用、权限体系和商业化版本。",
    )
    add_table(
        doc,
        ["阶段", "周期", "目标", "关键交付"],
        [
            ("Phase 1", "0-3 个月", "验证内容矩阵和核心工具闭环。", "10 个种子账号、爆款逆向器 MVP、角色圣经模板、100-300 条内容测试。"),
            ("Phase 2", "3-6 个月", "接入真实数据与多模态生成流程。", "ASR/OCR/抽帧能力、平台数据回流、内容日历、账号复盘报告。"),
            ("Phase 3", "6-12 个月", "开启 SaaS MVP 收费试点。", "Pro/Studio 版本、账号授权、协作审批、API/Agent 额度体系。"),
            ("Phase 4", "12 个月后", "形成内容 IP 与产品平台双轮增长。", "多账号矩阵、品牌共创、数据模型优化、商业化规模验证。"),
        ],
        [1.0, 1.0, 2.0, 2.4],
    )

    add_heading(doc, "八、商业模式", 1)
    add_para(
        doc,
        "业务采用双引擎收入结构：内容侧先建立现金流、案例和分发能力；产品侧承接更高毛利和可规模化收入；服务侧为品牌和 MCN 提供定制化账号孵化与内容矩阵解决方案。",
    )
    add_table(
        doc,
        ["收入类型", "收入来源", "阶段重点"],
        [
            ("内容商业化", "广告植入、IP 授权、周边、电商佣金、品牌共创。", "0-12 个月作为案例验证和现金流补充。"),
            ("产品订阅", "CreatorOS Pro / Studio / API：趋势监控、Agent 额度、多账号协作。", "6-12 个月开始收费试点，后续成为规模化重点。"),
            ("服务与共创", "品牌账号孵化、内容矩阵搭建、数据复盘、定制 Agent。", "早期用来获取 B 端需求和高质量案例。"),
        ],
        [1.4, 2.85, 2.1],
    )
    add_para(
        doc,
        "收入结构预期会随阶段变化：0-6 个月以内容商业化和服务为主；6-12 个月开始引入产品订阅；12-24 个月提高产品订阅和 API/Agent 收入占比。",
    )

    add_heading(doc, "九、市场进入策略", 1)
    add_para(
        doc,
        "市场进入策略分两条线推进：第一条是自有 AI 萌宠内容矩阵，用内容表现证明方法有效；第二条是 CreatorOS 产品化，把经过验证的流程提供给创作者、MCN 和品牌团队。",
    )
    add_bullet(doc, "种子账号策略：优先孵化 6 个方向、10 个种子账号，用统一角色资产库和内容流水线提高复用效率。")
    add_bullet(doc, "内容测试策略：每周围绕 3-5 个热点信号生成多版本内容，记录标题、钩子、完播、互动和转粉表现。")
    add_bullet(doc, "获客策略：通过 AI 工具导航站、创作者社群、内容创业训练营、案例挑战赛获取早期用户。")
    add_bullet(doc, "B 端策略：为品牌与 MCN 提供账号矩阵共创服务，沉淀可复制行业模板。")

    add_heading(doc, "十、组织与资源配置", 1)
    add_para(
        doc,
        "建议早期团队保持轻量，围绕产品、内容、模型工程和商务四类关键能力组建。外包可用于部分设计、剪辑和素材生产，但账号方法论、角色资产和数据复盘应掌握在核心团队内。",
    )
    add_table(
        doc,
        ["角色", "职责", "优先级"],
        [
            ("产品/增长负责人", "定义 CreatorOS 产品路线、用户旅程、指标体系和商业化策略。", "最高"),
            ("AI 工程/全栈", "实现爆款逆向、内容 Agent、多模态生成和平台数据接入。", "最高"),
            ("内容导演/编剧", "维护角色圣经、栏目库、脚本质量和账号内容调性。", "最高"),
            ("运营/投放", "账号发布、数据复盘、社区获客、内容测试。", "高"),
            ("商务合作", "品牌共创、MCN 试点、渠道合作和融资材料维护。", "中高"),
        ],
        [1.55, 3.8, 1.0],
    )

    add_heading(doc, "十一、融资计划与资金用途", 1)
    add_para(
        doc,
        "建议天使轮融资人民币 500 万，用于 12 个月验证“账号矩阵 + 产品化”双引擎。该金额是当前计划书建议口径，可根据投资人类型、估值预期和团队资源调整为 300 万、500 万或 1000 万三个版本。",
    )
    add_table(
        doc,
        ["用途", "比例", "资金方向"],
        [
            ("AI 内容生产与产品研发", "35%", "爆款逆向、内容 Agent、多模态工作流、SaaS MVP。"),
            ("账号矩阵运营与投放测试", "25%", "种子账号运营、内容测试、平台投放、账号复盘。"),
            ("内容团队与角色资产建设", "20%", "角色圣经、栏目库、脚本、分镜、素材资产。"),
            ("数据接入、版权与合规", "12%", "平台数据、授权内容、版权流程、内容审核。"),
            ("商务拓展与品牌合作", "8%", "品牌共创、MCN 试点、渠道合作、活动成本。"),
        ],
        [2.1, 0.9, 3.35],
    )
    add_callout(
        doc,
        "12 个月里程碑",
        "完成 10 个种子账号、覆盖 6 个萌宠+方向、累计 1000 条内容测试，并启动 CreatorOS Pro / Studio 付费试点。",
        fill=WARN_FILL,
    )

    add_heading(doc, "十二、关键指标", 1)
    add_para(doc, "建议用内容指标、产品指标和商业指标三层来管理早期验证，避免只看播放量而忽略可复用能力。")
    add_table(
        doc,
        ["指标类型", "核心指标", "验证意义"],
        [
            ("内容增长", "账号数、内容条数、播放/完播、互动率、转粉率、爆款率。", "验证萌宠+方向和内容生产方法是否有效。"),
            ("产品激活", "完成定位比例、生成 prompt 包比例、生成草稿比例、保存草稿比例。", "验证 CreatorOS 是否能降低创作者工作成本。"),
            ("留存复用", "7 日/30 日回访、每用户生成次数、监控赛道数、复盘次数。", "验证工具是否成为工作流而非一次性玩具。"),
            ("商业化", "付费转化、ARPU、品牌合作数、服务试点数、电商/授权收入。", "验证内容侧和产品侧收入潜力。"),
        ],
        [1.25, 2.65, 2.45],
    )

    add_heading(doc, "十三、风险与控制", 1)
    add_table(
        doc,
        ["风险", "影响", "控制方式"],
        [
            ("平台数据合规", "未经授权抓取或自动发布可能触发平台风险。", "优先使用开放 API、第三方授权数据、用户授权数据；保留人工确认。"),
            ("版权与同质化", "直接复刻热点内容会带来版权和品牌风险。", "爆款逆向器只复用结构和节奏，必须改写角色、台词、场景和 CTA。"),
            ("AI 内容质量", "批量生成可能出现低质、事实错误或风格漂移。", "建立角色圣经、审核机制、质量评分和复盘规则。"),
            ("模型与素材成本", "多模态生成成本可能影响毛利。", "分层生成策略，高价值内容使用高质量模型，常规内容使用模板化流程。"),
            ("商业化节奏", "账号流量和 SaaS 付费转化存在不确定性。", "内容侧、产品侧、服务侧并行验证，避免单一收入依赖。"),
        ],
        [1.55, 2.0, 2.8],
    )

    add_heading(doc, "十四、下一步行动计划", 1)
    add_number(doc, "完善 10 个种子账号的人设、栏目、封面规范和商业化路径。")
    add_number(doc, "将爆款逆向器从文本 MVP 升级到视频上传、抽帧、ASR、OCR、评论抓取的多模态版本。")
    add_number(doc, "建立 30 天内容测试日历，每周复盘不同“萌宠+”方向的表现。")
    add_number(doc, "用首批测试数据更新融资材料，将市场假设换成自有账号表现和客户访谈证据。")
    add_number(doc, "推进 CreatorOS Pro / Studio 试点，寻找 10-20 个愿意付费或深度共创的创作者/MCN/品牌团队。")

    add_heading(doc, "附录：资料来源与待核验项", 1)
    add_para(
        doc,
        "本业务计划书基于已完成的 CreatorOS AI 萌宠+ 内容工厂融资路演 PPT、CreatorOS PRD 和当前 MVP 原型整理。以下公开资料为 PPT 引用口径，正式融资版本建议补充报告原文、发布日期、链接截屏和数据授权说明。",
    )
    add_bullet(doc, "Financial Times：中国宠物消费及宠物经济相关报道。")
    add_bullet(doc, "Goldman Sachs：创作者经济 2027 年市场规模预测。")
    add_bullet(doc, "TechRadar / Adobe：创作者使用生成式 AI 的调研报道。")
    add_bullet(doc, "Washington Post：中国 AI 玩具/陪伴宠物市场相关报道。")
    add_bullet(doc, "CreatorOS 本地 MVP：趋势雷达、爆款逆向器、定位策略、内容 Agent、审批发布队列。")

    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUT)


if __name__ == "__main__":
    build_doc()
    print(OUT)
